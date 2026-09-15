
import json
import traceback
from io import BytesIO
from pathlib import Path
from datetime import datetime

import numpy as np
import pandas as pd
import streamlit as st

import bet_builder_v8_1_robust as engine


# ============================================================
# CONFIG
# ============================================================

APP_VERSION = "V8.1 ROBUST MOBILE — APUESTAS CLARAS"
DATA_DIR = Path("app_data")
DATA_DIR.mkdir(exist_ok=True)

META_FILE = DATA_DIR / "last_run.json"
TOP_FILE = DATA_DIR / "latest_top.csv"
VAR_FILE = DATA_DIR / "latest_variants.csv"
ALL_FILE = DATA_DIR / "latest_all.csv"

AUTO_REFRESH_HOURS = 12

st.set_page_config(
    page_title="Bet Builder V8.1",
    page_icon="⚽",
    layout="wide",
    initial_sidebar_state="collapsed",
)


# ============================================================
# CSS
# ============================================================

st.markdown(
    """
    <style>
      .block-container {
        padding-top: .7rem;
        padding-bottom: 3rem;
        max-width: 1080px;
      }

      .hero {
        padding: 1rem 1.05rem;
        border-radius: 20px;
        border: 1px solid rgba(128,128,128,.18);
        background: linear-gradient(135deg, rgba(16,185,129,.13), rgba(59,130,246,.09));
        margin-bottom: .7rem;
      }

      .hero h1 {
        margin: 0;
        font-size: clamp(1.55rem, 6vw, 2.25rem);
      }

      .hero p {
        margin: .4rem 0 0 0;
        opacity: .78;
      }

      .pill {
        display: inline-block;
        padding: .26rem .6rem;
        border-radius: 999px;
        font-weight: 800;
        font-size: .77rem;
        margin-right: .3rem;
        margin-top: .35rem;
      }

      .green { background: rgba(34,197,94,.17); color: #16a34a; }
      .amber { background: rgba(245,158,11,.18); color: #d97706; }
      .blue  { background: rgba(59,130,246,.16); color: #2563eb; }
      .gray  { background: rgba(107,114,128,.16); color: #6b7280; }
      .red   { background: rgba(239,68,68,.15); color: #dc2626; }

      .section-head {
        font-weight: 900;
        font-size: 1.22rem;
        margin: .9rem 0 .35rem 0;
      }

      .ticket {
        border: 1px solid rgba(128,128,128,.2);
        border-radius: 18px;
        padding: .95rem 1rem;
        margin: .7rem 0;
        background: rgba(128,128,128,.035);
      }

      .ticket-green {
        border-left: 7px solid #22c55e;
      }

      .ticket-amber {
        border-left: 7px solid #f59e0b;
      }

      .match {
        font-weight: 900;
        font-size: 1.08rem;
        margin-top: .35rem;
      }

      .sub {
        opacity: .70;
        font-size: .84rem;
        margin-bottom: .7rem;
      }

      .betline {
        display: grid;
        grid-template-columns: 36% 64%;
        border-top: 1px solid rgba(128,128,128,.12);
        padding: .45rem 0;
        gap: .4rem;
      }

      .betline b {
        opacity: .70;
        font-size: .82rem;
      }

      .betline span {
        font-weight: 800;
      }

      .numbers {
        display: grid;
        grid-template-columns: repeat(3, minmax(0,1fr));
        gap: .45rem;
        margin-top: .75rem;
      }

      .numbers > div {
        background: rgba(128,128,128,.06);
        border-radius: 12px;
        padding: .5rem .55rem;
      }

      .numbers b {
        display: block;
        opacity: .62;
        font-size: .72rem;
      }

      .numbers span {
        font-size: .98rem;
        font-weight: 900;
      }

      .instruction {
        padding: .75rem .85rem;
        border-radius: 14px;
        background: rgba(59,130,246,.09);
        border: 1px solid rgba(59,130,246,.18);
        margin-bottom: .55rem;
      }

      .stButton button, .stDownloadButton button {
        min-height: 46px;
        border-radius: 13px;
        font-weight: 800;
      }

      @media (max-width: 600px) {
        .betline {
          grid-template-columns: 39% 61%;
        }
        .numbers {
          grid-template-columns: repeat(3, minmax(0,1fr));
        }
      }
    </style>
    """,
    unsafe_allow_html=True,
)


# ============================================================
# HELPERS
# ============================================================

def safe_float(v, default=np.nan):
    try:
        x = float(v)
        return x if np.isfinite(x) else default
    except Exception:
        return default


def pct(v):
    x = safe_float(v)
    return "—" if pd.isna(x) else f"{x*100:.1f}%"


def num(v, digits=2):
    x = safe_float(v)
    return "—" if pd.isna(x) else f"{x:.{digits}f}"


def load_meta():
    if not META_FILE.exists():
        return {}
    try:
        return json.loads(META_FILE.read_text(encoding="utf-8"))
    except Exception:
        return {}


def save_meta(meta):
    META_FILE.write_text(
        json.dumps(meta, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def load_csv(path):
    if not path.exists():
        return pd.DataFrame()

    try:
        df = pd.read_csv(path)

        if "Fecha" in df.columns:
            df["Fecha"] = pd.to_datetime(
                df["Fecha"],
                errors="coerce",
            )

        return df

    except Exception:
        return pd.DataFrame()


def cache_is_current(meta):
    if not meta:
        return False

    try:
        now = datetime.now(engine.TZ_PERU)
        end = pd.Timestamp(meta["window_end"]).date()
        generated = datetime.fromisoformat(meta["generated_at"])

        if generated.tzinfo is None:
            return False

        age_hours = (
            now - generated
        ).total_seconds() / 3600

        return (
            now.date() <= end
            and age_hours <= AUTO_REFRESH_HOURS
        )

    except Exception:
        return False


# ============================================================
# MOTOR
# ============================================================

def run_model():
    inicio, fin = engine.ventana_objetivo()

    hist = engine.descargar_historico_total()
    ds, estados = engine.construir_dataset_v81(hist)

    schedule_index = engine.construir_schedule_index(hist)
    perfiles = engine.construir_perfil_arbitros(hist)
    context_h1 = engine.descargar_contexto_h1()

    fixtures = engine.descargar_fixtures_objetivo(
        inicio,
        fin,
    )

    resultado, variantes = engine.pronosticar_v81(
        fixtures,
        ds,
        estados,
        perfiles,
        schedule_index,
        context_h1,
    )

    top = engine.crear_top30_v81(resultado)

    top.to_csv(
        TOP_FILE,
        index=False,
        encoding="utf-8-sig",
    )

    variantes.to_csv(
        VAR_FILE,
        index=False,
        encoding="utf-8-sig",
    )

    resultado.to_csv(
        ALL_FILE,
        index=False,
        encoding="utf-8-sig",
    )

    meta = {
        "generated_at": datetime.now(
            engine.TZ_PERU
        ).isoformat(),
        "window_start": str(
            pd.Timestamp(inicio).date()
        ),
        "window_end": str(
            pd.Timestamp(fin).date()
        ),
        "n_top": int(len(top)),
        "n_bets": (
            int(
                (
                    top["Accion"]
                    == "APOSTAR"
                ).sum()
            )
            if not top.empty
            else 0
        ),
        "n_watch": (
            int(
                (
                    top["Accion"]
                    == "VIGILAR"
                ).sum()
            )
            if not top.empty
            else 0
        ),
        "version": APP_VERSION,
    }

    save_meta(meta)

    return (
        top,
        variantes,
        resultado,
        meta,
    )


def maybe_auto_generate():
    meta = load_meta()

    if cache_is_current(meta):
        return (
            load_csv(TOP_FILE),
            load_csv(VAR_FILE),
            load_csv(ALL_FILE),
            meta,
            False,
        )

    with st.status(
        "Actualizando la próxima ventana de 7 días…",
        expanded=True,
    ) as status:
        st.write(
            "Descargando histórico y calendario…"
        )

        top, var, all_df, meta = run_model()

        status.update(
            label="Pronósticos actualizados",
            state="complete",
            expanded=False,
        )

    return (
        top,
        var,
        all_df,
        meta,
        True,
    )


# ============================================================
# APUESTAS CLARAS
# ============================================================

def label_side(team, local, visita):
    if team == local:
        return f"{team} (LOCAL)"
    if team == visita:
        return f"{team} (VISITA)"
    return team


def mercados_de_fila(row):
    variante = str(
        row.get(
            "VarianteMasSegura",
            "BASE",
        )
        or "BASE"
    )

    local = str(
        row.get("Local", "")
        or ""
    )

    visita = str(
        row.get("Visitante", "")
        or ""
    )

    favorito = str(
        row.get("Favorito", "")
        or ""
    )

    dog = str(
        row.get("Underdog", "")
        or ""
    )

    fav = label_side(
        favorito,
        local,
        visita,
    )

    underdog = label_side(
        dog,
        local,
        visita,
    )

    resultado = "NO INCLUIDO"

    gol_equipo = (
        f"{underdog}: MARCA 1+ GOL"
    )

    corners_equipo = (
        f"{fav}: 4+ CÓRNERS"
    )

    total_goles = (
        "MENOS DE 4.5 GOLES"
    )

    tarjetas = "—"
    corners_1t = "—"
    gol_1t = "—"

    upper = variante.upper()

    if "BLINDADO" in upper:
        corners_equipo = (
            f"{fav}: 3+ CÓRNERS"
        )
        total_goles = (
            "MENOS DE 5.5 GOLES"
        )

    if "TARJETA" in upper:
        tarjetas = (
            "4+ TARJETAS TOTALES"
        )

    if (
        "CORNERS 1T" in upper
        or "CÓRNERS 1T" in upper
    ):
        corners_1t = (
            "3+ CÓRNERS TOTALES 1T"
        )

    if "GOL 1T" in upper:
        gol_1t = (
            "MÁS DE 0.5 GOLES 1T"
        )

    patas = [
        gol_equipo,
        corners_equipo,
        total_goles,
    ]

    for extra in (
        tarjetas,
        corners_1t,
        gol_1t,
    ):
        if extra != "—":
            patas.append(extra)

    return {
        "RESULTADO": resultado,
        "GOL DE EQUIPO": gol_equipo,
        "CÓRNERS DE EQUIPO": corners_equipo,
        "TOTAL GOLES": total_goles,
        "TARJETAS": tarjetas,
        "CÓRNERS 1T": corners_1t,
        "GOL 1T": gol_1t,
        "APUESTA COMPLETA": " + ".join(patas),
    }


def crear_apuestas_claras(top):
    columnas = [
        "N°",
        "ESTADO",
        "FECHA",
        "HORA",
        "COMPETICIÓN",
        "PARTIDO",
        "RESULTADO",
        "GOL DE EQUIPO",
        "CÓRNERS DE EQUIPO",
        "TOTAL GOLES",
        "TARJETAS",
        "CÓRNERS 1T",
        "GOL 1T",
        "CUOTA MÍNIMA",
        "CUOTA REAL",
        "CUMPLE CUOTA",
        "FIABILIDAD",
        "PROB. BUILDER",
        "APUESTA COMPLETA",
    ]

    if top is None or top.empty:
        return pd.DataFrame(
            columns=columnas
        )

    rows = []

    for i, (_, r) in enumerate(
        top.iterrows(),
        start=1,
    ):
        mercados = mercados_de_fila(
            r
        )

        fecha = r.get(
            "Fecha",
            "",
        )

        try:
            fecha = pd.Timestamp(
                fecha
            ).strftime(
                "%d/%m/%Y"
            )
        except Exception:
            pass

        rows.append({
            "N°": int(
                r.get(
                    "Ranking",
                    i,
                )
            ),
            "ESTADO": str(
                r.get(
                    "Accion",
                    "",
                )
            ),
            "FECHA": fecha,
            "HORA": str(
                r.get(
                    "HoraPeru",
                    "",
                )
            ),
            "COMPETICIÓN": str(
                r.get(
                    "Competicion",
                    "",
                )
            ),
            "PARTIDO": (
                f"{r.get('Local','')} "
                f"vs {r.get('Visitante','')}"
            ),
            **mercados,
            "CUOTA MÍNIMA": safe_float(
                r.get(
                    "CuotaMinMasSegura",
                    np.nan,
                )
            ),
            "CUOTA REAL": "",
            "CUMPLE CUOTA": "PENDIENTE",
            "FIABILIDAD": safe_float(
                r.get(
                    "ReliabilityScore",
                    np.nan,
                )
            ),
            "PROB. BUILDER": safe_float(
                r.get(
                    "P_VarianteMasSegura",
                    np.nan,
                )
            ),
        })

    return pd.DataFrame(
        rows,
        columns=columnas,
    )


# ============================================================
# EXCEL CLARO + TÉCNICO — GENERADO POR LA APP
# ============================================================

def excel_claro_bytes(
    top,
    variants,
    all_df,
):
    claras = crear_apuestas_claras(
        top
    )

    output = BytesIO()

    with pd.ExcelWriter(
        output,
        engine="openpyxl",
    ) as writer:
        claras.to_excel(
            writer,
            sheet_name="APUESTAS_CLARAS",
            index=False,
            startrow=4,
        )

        top.to_excel(
            writer,
            sheet_name="TOP_30_MAX",
            index=False,
        )

        if not variants.empty:
            variants.to_excel(
                writer,
                sheet_name="VARIANTES",
                index=False,
            )

        if not all_df.empty:
            all_df.to_excel(
                writer,
                sheet_name="TECNICO",
                index=False,
            )

        from openpyxl.styles import (
            Font,
            PatternFill,
            Alignment,
            Border,
            Side,
        )
        from openpyxl.utils import (
            get_column_letter,
        )

        wb = writer.book
        ws = wb[
            "APUESTAS_CLARAS"
        ]

        last_col = 19

        ws.merge_cells(
            start_row=1,
            start_column=1,
            end_row=1,
            end_column=last_col,
        )

        ws["A1"] = (
            "BET BUILDER V8.1 — APUESTAS CLARAS"
        )

        ws["A1"].font = Font(
            bold=True,
            color="FFFFFF",
            size=18,
        )

        ws["A1"].fill = PatternFill(
            "solid",
            fgColor="102A43",
        )

        ws["A1"].alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        ws.row_dimensions[1].height = 31

        ws.merge_cells(
            start_row=2,
            start_column=1,
            end_row=2,
            end_column=last_col,
        )

        ws["A2"] = (
            "PRIMERO LEE ESTA HOJA. Verde = candidato APOSTAR; "
            "amarillo = VIGILAR (NO apostar todavía). "
            "Escribe la cuota ofrecida por la casa en CUOTA REAL."
        )

        ws["A2"].fill = PatternFill(
            "solid",
            fgColor="D9EAF7",
        )

        ws["A2"].font = Font(
            bold=True,
            color="203040",
        )

        ws["A2"].alignment = Alignment(
            wrap_text=True,
            vertical="center",
        )

        ws.row_dimensions[2].height = 36

        ws["A3"] = "🟢 APOSTAR"
        ws["B3"] = "🟡 VIGILAR"
        ws["C3"] = (
            "RESULTADO/DOBLE OPORTUNIDAD = NO INCLUIDO "
            "porque el modelo actual no lo calcula."
        )

        ws["A3"].fill = PatternFill(
            "solid",
            fgColor="C6EFCE",
        )

        ws["B3"].fill = PatternFill(
            "solid",
            fgColor="FFF2CC",
        )

        ws["C3"].fill = PatternFill(
            "solid",
            fgColor="F2F2F2",
        )

        for c in (
            "A3",
            "B3",
            "C3",
        ):
            ws[c].font = Font(
                bold=True
            )
            ws[c].alignment = Alignment(
                wrap_text=True,
            )

        header_row = 5

        headers = {
            cell.value: cell.column
            for cell in ws[
                header_row
            ]
        }

        for cell in ws[
            header_row
        ]:
            cell.font = Font(
                bold=True,
                color="FFFFFF",
            )
            cell.fill = PatternFill(
                "solid",
                fgColor="1F4E78",
            )
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        ws.freeze_panes = "A6"

        widths = {
            "A": 5,
            "B": 12,
            "C": 13,
            "D": 8,
            "E": 22,
            "F": 29,
            "G": 17,
            "H": 34,
            "I": 34,
            "J": 20,
            "K": 23,
            "L": 24,
            "M": 21,
            "N": 14,
            "O": 14,
            "P": 15,
            "Q": 13,
            "R": 14,
            "S": 65,
        }

        for col, width in widths.items():
            ws.column_dimensions[
                col
            ].width = width

        thin = Side(
            style="thin",
            color="D9E2F3",
        )

        for rr in range(
            6,
            ws.max_row + 1,
        ):
            estado = str(
                ws.cell(
                    rr,
                    headers[
                        "ESTADO"
                    ],
                ).value
                or ""
            )

            fill = (
                "E2F0D9"
                if estado
                == "APOSTAR"
                else "FFF2CC"
            )

            for cc in range(
                1,
                last_col + 1,
            ):
                cell = ws.cell(
                    rr,
                    cc,
                )
                cell.fill = PatternFill(
                    "solid",
                    fgColor=fill,
                )
                cell.alignment = Alignment(
                    wrap_text=True,
                    vertical="top",
                )
                cell.border = Border(
                    bottom=thin,
                )

            ws.row_dimensions[
                rr
            ].height = 54

            # Editable cuota real.
            real_cell = ws.cell(
                rr,
                headers[
                    "CUOTA REAL"
                ],
            )
            real_cell.fill = PatternFill(
                "solid",
                fgColor="FFFFFF",
            )
            real_cell.number_format = (
                "0.00"
            )

            min_cell = ws.cell(
                rr,
                headers[
                    "CUOTA MÍNIMA"
                ],
            )
            min_cell.fill = PatternFill(
                "solid",
                fgColor="D9EAF7",
            )
            min_cell.font = Font(
                bold=True,
                color="1F4E78",
            )
            min_cell.number_format = (
                "0.00"
            )

            # Fórmula SI/NO.
            cr = get_column_letter(
                headers[
                    "CUOTA REAL"
                ]
            )
            cm = get_column_letter(
                headers[
                    "CUOTA MÍNIMA"
                ]
            )

            ws.cell(
                rr,
                headers[
                    "CUMPLE CUOTA"
                ],
            ).value = (
                f'=IF({cr}{rr}="","PENDIENTE",'
                f'IF(AND({cr}{rr}>={cm}{rr},'
                f'{cr}{rr}>=4.20),"SI","NO"))'
            )

            for h in (
                "FIABILIDAD",
                "PROB. BUILDER",
            ):
                ws.cell(
                    rr,
                    headers[h],
                ).number_format = (
                    "0.0%"
                )

            full = ws.cell(
                rr,
                headers[
                    "APUESTA COMPLETA"
                ],
            )

            full.fill = PatternFill(
                "solid",
                fgColor="EAF2F8",
            )

            full.font = Font(
                bold=True,
                color="17365D",
            )

        # Hojas técnicas con cabecera estándar.
        for name in wb.sheetnames:
            if name == "APUESTAS_CLARAS":
                continue

            sh = wb[
                name
            ]

            if sh.max_row >= 1:
                for cell in sh[1]:
                    cell.font = Font(
                        bold=True,
                        color="FFFFFF",
                    )
                    cell.fill = PatternFill(
                        "solid",
                        fgColor="1F4E78",
                    )
                    cell.alignment = Alignment(
                        horizontal="center",
                        wrap_text=True,
                    )

                sh.freeze_panes = "A2"
                sh.auto_filter.ref = (
                    sh.dimensions
                )

    output.seek(0)

    return output.getvalue()


# ============================================================
# CARD APUESTA CLARA
# ============================================================

def action_class(action):
    return (
        "green"
        if action == "APOSTAR"
        else "amber"
    )


def render_apuesta_clara(
    row,
    idx,
):
    action = str(
        row.get(
            "Accion",
            "VIGILAR",
        )
    )

    markets = mercados_de_fila(
        row
    )

    local = str(
        row.get(
            "Local",
            "",
        )
    )

    visita = str(
        row.get(
            "Visitante",
            "",
        )
    )

    comp = str(
        row.get(
            "Competicion",
            "",
        )
    )

    fecha = row.get(
        "Fecha",
        "",
    )

    try:
        fecha = pd.Timestamp(
            fecha
        ).strftime(
            "%d/%m"
        )
    except Exception:
        fecha = str(
            fecha
        )

    hora = str(
        row.get(
            "HoraPeru",
            "",
        )
    )

    qmin = safe_float(
        row.get(
            "CuotaMinMasSegura",
            np.nan,
        )
    )

    rel = safe_float(
        row.get(
            "ReliabilityScore",
            np.nan,
        )
    )

    pb = safe_float(
        row.get(
            "P_VarianteMasSegura",
            np.nan,
        )
    )

    border = (
        "ticket-green"
        if action == "APOSTAR"
        else "ticket-amber"
    )

    st.markdown(
        f"""
        <div class="ticket {border}">
          <div>
            <span class="pill {action_class(action)}">{action}</span>
            <span class="pill blue">{comp}</span>
          </div>

          <div class="match">{local} vs {visita}</div>
          <div class="sub">{fecha} · {hora} PET</div>

          <div class="betline"><b>RESULTADO</b><span>{markets['RESULTADO']}</span></div>
          <div class="betline"><b>GOL DE EQUIPO</b><span>{markets['GOL DE EQUIPO']}</span></div>
          <div class="betline"><b>CÓRNERS EQUIPO</b><span>{markets['CÓRNERS DE EQUIPO']}</span></div>
          <div class="betline"><b>TOTAL GOLES</b><span>{markets['TOTAL GOLES']}</span></div>
          <div class="betline"><b>TARJETAS</b><span>{markets['TARJETAS']}</span></div>
          <div class="betline"><b>CÓRNERS 1T</b><span>{markets['CÓRNERS 1T']}</span></div>
          <div class="betline"><b>GOL 1T</b><span>{markets['GOL 1T']}</span></div>

          <div class="numbers">
            <div><b>CUOTA MÍN.</b><span>{num(qmin)}</span></div>
            <div><b>FIABILIDAD</b><span>{pct(rel)}</span></div>
            <div><b>PROB.</b><span>{pct(pb)}</span></div>
          </div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    quota_real = st.number_input(
        "Cuota real ofrecida por la casa",
        min_value=0.0,
        max_value=50.0,
        value=0.0,
        step=0.05,
        key=f"quota_real_{idx}",
    )

    if action != "APOSTAR":
        st.warning(
            "VIGILAR: no apostar todavía aunque la cuota sea alta. "
            "El partido no superó todos los filtros estrictos."
        )

    elif quota_real <= 0:
        st.info(
            f"Escribe la cuota real. Debe ser ≥ {qmin:.2f} "
            "y nunca inferior a 4.20."
        )

    elif quota_real >= max(
        qmin,
        4.20,
    ):
        ev = (
            pb * quota_real - 1
            if pd.notna(
                pb
            )
            else np.nan
        )

        st.success(
            "✅ CUOTA VÁLIDA"
            + (
                f" · EV modelo ≈ {ev*100:.1f}%"
                if pd.notna(
                    ev
                )
                else ""
            )
        )

    else:
        st.error(
            f"❌ NO TOMAR A ESA CUOTA. "
            f"Se exige ≥ {max(qmin,4.20):.2f}."
        )

    with st.expander(
        "Ver apuesta completa"
    ):
        st.write(
            markets[
                "APUESTA COMPLETA"
            ]
        )


# ============================================================
# HEADER
# ============================================================

st.markdown(
    """
    <div class="hero">
      <h1>⚽ Bet Builder V8.1</h1>
      <p><b>APUESTAS CLARAS primero</b> · después los datos técnicos</p>
      <span class="pill green">APOSTAR</span>
      <span class="pill amber">VIGILAR</span>
      <span class="pill blue">CUOTA ≥ 4.20</span>
    </div>
    """,
    unsafe_allow_html=True,
)


# ============================================================
# LOAD
# ============================================================

try:
    (
        top,
        variants,
        all_df,
        meta,
        refreshed,
    ) = maybe_auto_generate()

except Exception as e:
    st.error(
        "No fue posible actualizar el modelo."
    )
    st.code(
        str(e)
    )

    with st.expander(
        "Detalle técnico"
    ):
        st.code(
            traceback.format_exc()
        )

    top = load_csv(
        TOP_FILE
    )
    variants = load_csv(
        VAR_FILE
    )
    all_df = load_csv(
        ALL_FILE
    )
    meta = load_meta()


# ============================================================
# CONTROLES
# ============================================================

c1, c2 = st.columns(
    2
)

with c1:
    if st.button(
        "🔄 Actualizar ahora",
        use_container_width=True,
        type="primary",
    ):
        try:
            with st.status(
                "Ejecutando V8.1 ROBUST…",
                expanded=True,
            ) as status:
                (
                    top,
                    variants,
                    all_df,
                    meta,
                ) = run_model()

                status.update(
                    label="Actualización terminada",
                    state="complete",
                    expanded=False,
                )

            st.rerun()

        except Exception as e:
            st.error(
                str(e)
            )

with c2:
    if not top.empty:
        excel_bytes = excel_claro_bytes(
            top,
            variants,
            all_df,
        )

        st.download_button(
            "📊 Excel APUESTAS CLARAS",
            data=excel_bytes,
            file_name=(
                "BET_BUILDER_V81_APUESTAS_CLARAS.xlsx"
            ),
            mime=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
            use_container_width=True,
        )


if meta:
    start = meta.get(
        "window_start",
        "—",
    )

    end = meta.get(
        "window_end",
        "—",
    )

    generated = meta.get(
        "generated_at",
        "",
    )

    try:
        gen_txt = datetime.fromisoformat(
            generated
        ).strftime(
            "%d/%m %H:%M"
        )
    except Exception:
        gen_txt = generated

    st.caption(
        f"Ventana: {start} → {end} · Actualizado: {gen_txt}"
    )


# ============================================================
# SECCIÓN PRINCIPAL — APUESTAS CLARAS
# ============================================================

st.markdown(
    '<div class="section-head">🎯 APUESTAS CLARAS</div>',
    unsafe_allow_html=True,
)

st.markdown(
    """
    <div class="instruction">
      <b>Cómo usar:</b> primero busca filas verdes <b>APOSTAR</b>.
      Construye exactamente los mercados que aparecen debajo.
      Los amarillos <b>VIGILAR</b> son informativos: no se toman todavía.
      El modelo actual no incluye 1X2/doble oportunidad, por eso RESULTADO
      muestra “NO INCLUIDO”.
    </div>
    """,
    unsafe_allow_html=True,
)

if top.empty:
    st.warning(
        "No hay candidatos modelables en esta ventana."
    )

else:
    n_bets = int(
        (
            top[
                "Accion"
            ]
            == "APOSTAR"
        ).sum()
    )

    n_watch = int(
        (
            top[
                "Accion"
            ]
            == "VIGILAR"
        ).sum()
    )

    a, b, c = st.columns(
        3
    )

    a.metric(
        "🟢 APOSTAR",
        n_bets,
    )

    b.metric(
        "🟡 VIGILAR",
        n_watch,
    )

    c.metric(
        "Oportunidades",
        len(
            top
        ),
    )

    view_option = st.radio(
        "Mostrar",
        [
            "APOSTAR primero",
            "Solo APOSTAR",
            "Solo VIGILAR",
            "Todos",
        ],
        horizontal=True,
    )

    view = top.copy()

    if view_option == "Solo APOSTAR":
        view = view[
            view[
                "Accion"
            ]
            == "APOSTAR"
        ]

    elif view_option == "Solo VIGILAR":
        view = view[
            view[
                "Accion"
            ]
            == "VIGILAR"
        ]

    elif view_option == "APOSTAR primero":
        pref = {
            "APOSTAR": 0,
            "VIGILAR": 1,
        }

        view[
            "_pref"
        ] = (
            view[
                "Accion"
            ]
            .map(
                pref
            )
            .fillna(
                9
            )
        )

        view = (
            view.sort_values(
                [
                    "_pref",
                    "ScoreRobusto",
                ],
                ascending=[
                    True,
                    False,
                ],
            )
            .drop(
                columns=[
                    "_pref"
                ],
                errors="ignore",
            )
        )

    if view.empty:
        st.info(
            "No hay partidos en esta categoría."
        )

    for idx, (_, row) in enumerate(
        view.iterrows(),
        start=1,
    ):
        render_apuesta_clara(
            row,
            idx,
        )


# ============================================================
# SECCIÓN TÉCNICA SECUNDARIA
# ============================================================

st.markdown(
    '<div class="section-head">📐 DATOS TÉCNICOS</div>',
    unsafe_allow_html=True,
)

with st.expander(
    "Resumen técnico"
):
    if not top.empty:
        avg_rel = pd.to_numeric(
            top[
                "ReliabilityScore"
            ],
            errors="coerce",
        ).mean()

        min_q = pd.to_numeric(
            top[
                "CuotaMinMasSegura"
            ],
            errors="coerce",
        ).min()

        x1, x2 = st.columns(
            2
        )

        x1.metric(
            "Fiabilidad media",
            pct(
                avg_rel
            ),
        )

        x2.metric(
            "Cuota mínima más baja",
            num(
                min_q
            ),
        )

with st.expander(
    "🧪 Variantes y probabilidades"
):
    if variants.empty:
        st.write(
            "No hay variantes calculadas."
        )

    else:
        cols = [
            c
            for c in [
                "Fecha",
                "Competicion",
                "Local",
                "Visitante",
                "AccionPartido",
                "Variante",
                "P_Conjunta",
                "CuotaJusta",
                "CuotaMinExigida",
                "Evidencia",
            ]
            if c in variants.columns
        ]

        st.dataframe(
            variants[
                cols
            ],
            use_container_width=True,
            hide_index=True,
        )


with st.expander(
    "ℹ️ Qué significa cada estado"
):
    st.markdown(
        """
        **APOSTAR:** superó los filtros estadísticos estrictos.
        Todavía hay que comprobar que la cuota real sea igual o superior
        a la cuota mínima.

        **VIGILAR:** tiene señales interesantes, pero no superó todos los
        requisitos. No debe tomarse como una recomendación de apuesta.

        **Fiabilidad:** calidad/contexto del modelo; no es la probabilidad
        de acertar.

        **Probabilidad:** estimación del builder completo elegido.
        """
    )


st.caption(
    "Modelo estadístico experimental. No garantiza ganancias. "
    "No se fuerza ninguna apuesta para completar 30."
)
